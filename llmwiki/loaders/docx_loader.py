"""Word (.docx) loader via python-docx."""

from __future__ import annotations

from pathlib import Path

from .base import LoaderError, LoadResult


def load(path: Path, **_kwargs) -> LoadResult:
    try:
        import docx  # python-docx
    except ImportError as e:  # pragma: no cover
        raise LoaderError("python-docx is required for .docx files.") from e

    document = docx.Document(str(path))
    lines: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style else ""
        if style.startswith("heading"):
            level = "".join(c for c in style if c.isdigit())
            hashes = "#" * (int(level) if level else 2)
            lines.append(f"{hashes} {text}")
        else:
            lines.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append("| " + " | ".join(cells) + " |")

    markdown = "\n\n".join(lines)
    if not markdown.strip():
        raise LoaderError(f"No extractable text in {path}.")
    return LoadResult(markdown=markdown, kind="docx", title=path.stem)
