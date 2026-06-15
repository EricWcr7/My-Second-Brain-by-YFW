"""PowerPoint (.pptx) loader via python-pptx."""

from __future__ import annotations

from pathlib import Path

from .base import LoaderError, LoadResult


def load(path: Path, **_kwargs) -> LoadResult:
    try:
        from pptx import Presentation  # python-pptx
    except ImportError as e:  # pragma: no cover
        raise LoaderError("python-pptx is required for .pptx files.") from e

    prs = Presentation(str(path))
    blocks: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs).strip()
                    if line:
                        texts.append(line)
        if texts:
            body = "\n".join(f"- {t}" for t in texts)
            blocks.append(f"## Slide {i}\n\n{body}")

    markdown = "\n\n".join(blocks)
    if not markdown.strip():
        raise LoaderError(f"No extractable text in {path}.")
    return LoadResult(markdown=markdown, kind="pptx", title=path.stem)
